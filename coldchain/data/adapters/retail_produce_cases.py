"""Adapter for published aggregate produce-display-case measurements."""

from collections.abc import Iterator

from openpyxl import load_workbook

from .base import DatasetAdapter
from ..schema import SensorObservation


class RetailProduceCasesAdapter(DatasetAdapter):
    dataset_id = "retail_produce_display_cases_v1"

    def iter_observations(self) -> Iterator[SensorObservation]:
        workbook = load_workbook(self.source, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() for value in next(rows)]
        for row_index, values in enumerate(rows, start=2):
            row = dict(zip(headers, values, strict=True))
            if row.get("MEANTEMP") is None:
                continue
            retailer = str(row.get("RETAILER", "unknown")).strip()
            case = str(row.get("CASE", "unknown")).strip()
            position = str(row.get("POSITION", "unknown")).strip()
            yield SensorObservation(
                dataset_id=self.dataset_id,
                shipment_id=f"{retailer}-case-{case}",
                sensor_id=position,
                observed_at=None,
                temperature_c=float(row["MEANTEMP"]),
                relative_humidity_pct=float(row["MEANMOIST"]) if row.get("MEANMOIST") is not None else None,
                product="fresh produce (case inventory not specified)",
                stage="retail refrigerated display",
                measurement_type="aggregate",
                source_row=f"Sheet1:{row_index}",
                metadata={
                    "doors": row.get("DOORS"),
                    "median_temperature_c": row.get("MEDIANTEMP"),
                    "minimum_temperature_c": row.get("MINTEMP"),
                    "maximum_temperature_c": row.get("MAXTEMP"),
                    "percent_above_source_limit": row.get("PABUSEHIGH"),
                    "longest_high_abuse_run": row.get("CONSECUTIVEABUSEHIGH"),
                    "timestamp_missing_by_source": True,
                },
            )

